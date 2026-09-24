"""Row count check: the XML counter and the DuckDB stage must agree, blanks excluded."""

import openpyxl
import pytest
from openpyxl.styles import Font

from etl import ingest, manifest
from etl.columns import RawFile
from etl.xlsx_count import count_rows, read_header

HEADER = ["CASE_NUMBER", "CASE_STATUS", "EMPLOYER_NAME"]


def _xlsx(path, rows, blank_styled_rows=0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    # Formatted but empty rows, like the ones DOL leaves at the bottom of its sheets.
    for i in range(blank_styled_rows):
        ws.cell(row=len(rows) + 2 + i, column=1).font = Font(bold=True)
    wb.save(path)


ROWS = [
    ["I-200-25001-000001", "Certified", "ACME LLC"],
    ["I-200-25001-000002", "Denied", "ACME LLC"],
    ["I-200-25001-000003", "Withdrawn", "Other Inc"],
]


def test_counter_skips_trailing_formatted_rows(tmp_path):
    p = tmp_path / "f.xlsx"
    _xlsx(p, ROWS, blank_styled_rows=25)
    c = count_rows(p)
    assert c.data_rows == 3
    assert c.blank_rows == 25
    assert c.first_blank_row == 5
    assert read_header(p) == HEADER


@pytest.fixture
def sandbox(tmp_path, monkeypatch):
    monkeypatch.setattr(manifest, "RAW", tmp_path)
    monkeypatch.setattr(manifest, "MANIFEST", tmp_path / "manifest.json")
    monkeypatch.setattr(ingest, "STAGED", tmp_path / "staged")
    monkeypatch.setattr(ingest, "ROOT", tmp_path)
    monkeypatch.setattr(
        ingest,
        "missing_columns",
        lambda kind, fy, q, header: [c for c in HEADER if c not in header],
    )
    return tmp_path


def test_stage_passes_and_records_counts(sandbox):
    p = sandbox / "LCA_Test.xlsx"
    _xlsx(p, ROWS, blank_styled_rows=10)
    rf = RawFile("LCA_Test.xlsx", "lca", 2025, 4, "https://example.test/f.xlsx", False)
    entry = ingest.stage_file(ingest.connect(":memory:"), rf, raw_dir=sandbox)
    assert entry["raw_rows"] == entry["loaded_rows"] == 3
    assert entry["blank_rows_in_sheet"] == 10
    assert len(entry["sha256"]) == 64


def test_stage_fails_loudly_on_count_mismatch(sandbox, monkeypatch):
    p = sandbox / "LCA_Test.xlsx"
    _xlsx(p, ROWS)
    real = count_rows

    def off_by_one(path):
        c = real(path)
        c.data_rows += 1
        return c

    monkeypatch.setattr(ingest, "count_rows", off_by_one)
    rf = RawFile("LCA_Test.xlsx", "lca", 2025, 4, "https://example.test/f.xlsx", False)
    with pytest.raises(ingest.RowCountError):
        ingest.stage_file(ingest.connect(":memory:"), rf, raw_dir=sandbox)


def test_stage_fails_on_missing_column(sandbox):
    p = sandbox / "LCA_Test.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["CASE_NUMBER", "CASE_STATUS"])
    wb.active.append(["I-200-25001-000001", "Certified"])
    wb.save(p)
    rf = RawFile("LCA_Test.xlsx", "lca", 2025, 4, "https://example.test/f.xlsx", False)
    with pytest.raises(KeyError):
        ingest.stage_file(ingest.connect(":memory:"), rf, raw_dir=sandbox)
